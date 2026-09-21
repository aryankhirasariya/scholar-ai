import io
from datetime import datetime
from typing import List, Dict, Optional

from docx import Document
from docx.shared import Pt, RGBColor

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib import colors


def generate_docx(title: Optional[str], sections: List[Dict]) -> bytes:
    """
    sections: list of {"question": str, "answer": str, "sources": list[str] | None}
    """
    doc = Document()

    doc.add_heading(title or "Scholar AI Report", level=0)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(f"Generated {datetime.now().strftime('%d %B %Y, %H:%M')}")
    meta_run.italic = True
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    for i, sec in enumerate(sections, 1):
        doc.add_heading(f"{i}. {sec.get('question', '').strip()}", level=2)
        doc.add_paragraph(sec.get("answer", "").strip())

        sources = sec.get("sources") or []
        if sources:
            src_para = doc.add_paragraph()
            run = src_para.add_run("Sources: " + ", ".join(sources))
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def generate_pdf(title: Optional[str], sections: List[Dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=22, spaceAfter=6
    )
    meta_style = ParagraphStyle(
        "Meta", parent=styles["Normal"], textColor=colors.HexColor("#777777"),
        fontSize=9, spaceAfter=18
    )
    q_style = ParagraphStyle(
        "Question", parent=styles["Heading2"], fontSize=13,
        spaceBefore=16, spaceAfter=6
    )
    a_style = ParagraphStyle(
        "Answer", parent=styles["Normal"], fontSize=10.5,
        leading=16, alignment=TA_LEFT
    )
    src_style = ParagraphStyle(
        "Sources", parent=styles["Normal"], fontSize=8.5,
        textColor=colors.HexColor("#888888"), spaceBefore=6
    )

    story = [
        Paragraph(title or "Scholar AI Report", title_style),
        Paragraph(f"Generated {datetime.now().strftime('%d %B %Y, %H:%M')}", meta_style),
    ]

    for i, sec in enumerate(sections, 1):
        question = (sec.get("question") or "").strip()
        answer = (sec.get("answer") or "").strip().replace("\n", "<br/>")

        story.append(Paragraph(f"{i}. {question}", q_style))
        story.append(Paragraph(answer, a_style))

        sources = sec.get("sources") or []
        if sources:
            story.append(Paragraph("Sources: " + ", ".join(sources), src_style))

        story.append(HRFlowable(width="100%", color=colors.HexColor("#dddddd"),
                                 spaceBefore=10, spaceAfter=4))

    doc.build(story)
    buf.seek(0)
    return buf.read()
